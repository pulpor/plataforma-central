import io
from datetime import date

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import SystemConfig

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_excel(config: SystemConfig, results: list[dict]) -> tuple[io.BytesIO, str]:
    """Generate a formatted .xlsx file for a system's results.

    Includes bold/filled header, autofilter, frozen header row, auto-sized
    columns and clickable hyperlinks for columns typed as "url".
    """
    columns = config.output.columns
    if columns:
        keys = [column.key for column in columns]
        labels = {column.key: column.label for column in columns}
        url_keys = {column.key for column in columns if column.type == "url"}
    else:
        keys = list(results[0].keys()) if results else []
        labels = {key: key for key in keys}
        url_keys = set()

    rows = [{key: record.get(key, "") for key in keys} for record in results]
    dataframe = pd.DataFrame(rows, columns=keys)
    dataframe.rename(columns=labels, inplace=True)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        dataframe.to_excel(writer, index=False, sheet_name="Resultados")
        worksheet = writer.sheets["Resultados"]

        for col_idx in range(1, len(keys) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        max_row = len(rows) + 1
        if keys:
            worksheet.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{max_row}"
        worksheet.freeze_panes = "A2"

        for col_idx, key in enumerate(keys, start=1):
            letter = get_column_letter(col_idx)
            max_len = len(str(labels.get(key, key)))
            for row in rows:
                max_len = max(max_len, len(str(row.get(key, ""))))
            worksheet.column_dimensions[letter].width = min(max_len + 4, 60)

            if key in url_keys:
                for row_idx, row in enumerate(rows, start=2):
                    url = row.get(key)
                    if url:
                        cell = worksheet.cell(row=row_idx, column=col_idx)
                        cell.hyperlink = url
                        cell.style = "Hyperlink"

    buffer.seek(0)
    prefix = config.export.filename_prefix or config.id
    filename = f"{prefix}-{date.today().isoformat()}.xlsx"
    return buffer, filename
